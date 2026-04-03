"""
从 Markdown 条目文件生成《中国大百科全书》风格 Word 文档。

依赖：
  pip install python-docx math2docx

用法：
  python create_entry.py -f 反向传播算法.md
  python create_entry.py -f 反向传播算法.md -o 反向传播算法.docx
  # 不传 -o 时，默认输出到当前文件夹（cwd）

Markdown 约定（受控子集）：
1. 可选 YAML front matter（使用 --- 包裹），支持字段：
   - title / 标题
   - pinyin / 拼音
   - english_name / 英文名称
   - alias / 又称
   - discipline / 所属学科
   - summary / 摘要
   - author / 作者
   - toc / 目录（可写为列表）
2. 正文标题使用 #/##/### ...
3. 行内公式使用 $...$
4. 独立公式使用 $$...$$（支持多行），公式编号自动按出现顺序生成。
5. 参考文献放在“参考文献”标题下，普通段落或列表项均可。
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

try:
    import math2docx  # LaTeX -> Word 原生公式（OMML）
except ImportError:
    math2docx = None

# 字体与尺寸（按样条图片视觉重构）
FONT_SONG = "宋体"
FONT_HEI = "黑体"
FONT_EN = "Times New Roman"

SIZE_TITLE = Pt(16)
SIZE_ABSTRACT = Pt(11)
SIZE_META = Pt(10.5)
SIZE_TOC = Pt(11)
SIZE_SECTION = Pt(18)
SIZE_SECTION_REF = Pt(13)
SIZE_BODY = Pt(11)
SIZE_FORMULA = Pt(11.5)
SIZE_REF = Pt(9.5)
SIZE_AUTHOR = Pt(9.5)

LINE_TITLE = 24
LINE_META = 18
LINE_BODY = 20
LINE_SECTION = 24
LINE_REF = 16

COLOR_LINE = "D9D9D9"
COLOR_BLUE = "2F80ED"

PAGE_TEXT_WIDTH_CM = 16.8  # A4(21cm) - 左右页边距各2.1cm

HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*$")
NUMBER_LINE_RE = re.compile(r"^[（(]\d+[）)]$")
FRONT_MATTER_ITEM_RE = re.compile(r"^\s*([A-Za-z0-9_\-\u4e00-\u9fff]+)\s*[:：]\s*(.*)$")
FRONT_MATTER_LIST_RE = re.compile(r"^\s*-\s+(.*)$")


@dataclass
class Block:
    kind: str  # "paragraph" | "equation"
    content: str


@dataclass
class Section:
    title: str
    blocks: list[Block]


@dataclass
class EntryData:
    title: str
    summary: str
    pinyin: str
    english_name: str
    alias: str
    discipline: str
    author: str
    toc: list[str]
    sections: list[Section]
    references: list[str]


def set_run_font(run, cn_font, en_font=FONT_EN, size=SIZE_BODY, bold=False, italic=False):
    run.font.name = en_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)
    rfonts.set(qn("w:ascii"), en_font)
    rfonts.set(qn("w:hAnsi"), en_font)
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic


def m(latex, fallback=None):
    """标记一个行内数学公式片段（优先用 Word 原生公式，失败时回退文本）。"""
    return {"latex": normalize_latex_for_word(latex), "fallback": fallback or latex}


def normalize_latex_for_word(latex: str) -> str:
    """
    对 LaTeX 做最小化归一化，匹配 Word 公式渲染特性：
    - 将 \\cdot 替换为更细的中点字符，避免 Word 中显示过粗过大。
    """
    return latex.replace(r"\cdot", "·")


def add_rich_content(para, content, cn_font=FONT_SONG, en_font=FONT_EN, size=SIZE_BODY, italic=False):
    """
    向段落写入混排内容：
    - str: 普通文本
    - {"latex": ..., "fallback": ...}: Word 行内数学公式
    """
    has_math = False
    if isinstance(content, str):
        content = [content]
    for seg in content:
        if isinstance(seg, str):
            r = para.add_run(seg)
            set_run_font(r, cn_font, en_font=en_font, size=size, italic=italic)
            continue
        if isinstance(seg, dict) and "latex" in seg:
            has_math = True
            if math2docx is not None:
                try:
                    math2docx.add_math(para, seg["latex"])
                    continue
                except Exception:
                    pass
            r = para.add_run(seg.get("fallback", seg["latex"]))
            set_run_font(r, cn_font, en_font=en_font, size=size, italic=italic)
            continue
        raise TypeError(f"Unsupported rich segment: {type(seg)!r}")
    return has_math


def set_para_spacing(para, before=0, after=0, line=None):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)


def set_paragraph_borders(paragraph, top=None, bottom=None, left=None, right=None):
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:pBdr"))
    if old is not None:
        ppr.remove(old)
    specs = {"top": top, "bottom": bottom, "left": left, "right": right}
    if not any(specs.values()):
        return
    pbdr = OxmlElement("w:pBdr")
    for edge, spec in specs.items():
        if not spec:
            continue
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), spec.get("val", "single"))
        if "sz" in spec:
            e.set(qn("w:sz"), str(spec["sz"]))
        if "space" in spec:
            e.set(qn("w:space"), str(spec["space"]))
        if "color" in spec:
            e.set(qn("w:color"), spec["color"])
        pbdr.append(e)
    ppr.append(pbdr)


def _set_border_xml(parent, tag_name, specs):
    old = parent.find(qn(tag_name))
    if old is not None:
        parent.remove(old)
    if not any(specs.values()):
        return
    borders = OxmlElement(tag_name)
    for edge, spec in specs.items():
        if not spec:
            continue
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), spec.get("val", "single"))
        if "sz" in spec:
            e.set(qn("w:sz"), str(spec["sz"]))
        if "space" in spec:
            e.set(qn("w:space"), str(spec["space"]))
        if "color" in spec:
            e.set(qn("w:color"), spec["color"])
        borders.append(e)
    parent.append(borders)


def set_table_borders(table, top=None, bottom=None, left=None, right=None, insideH=None, insideV=None):
    tblpr = table._tbl.tblPr
    specs = {
        "top": top,
        "bottom": bottom,
        "left": left,
        "right": right,
        "insideH": insideH,
        "insideV": insideV,
    }
    _set_border_xml(tblpr, "w:tblBorders", specs)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcpr = cell._tc.get_or_add_tcPr()
    specs = {"top": top, "bottom": bottom, "left": left, "right": right}
    _set_border_xml(tcpr, "w:tcBorders", specs)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.find(qn("w:tcMar"))
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for key, value in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
        el = tcmar.find(qn(f"w:{key}"))
        if el is None:
            el = OxmlElement(f"w:{key}")
            tcmar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_table_fixed_layout(table):
    tblpr = table._tbl.tblPr
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")


def line_spec(color=COLOR_LINE, sz=6, space=1):
    return {"val": "single", "sz": sz, "space": space, "color": color}


def add_entry_title(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_spacing(p, before=0, after=4, line=LINE_TITLE)
    r = p.add_run(text)
    set_run_font(r, FONT_SONG, size=SIZE_TITLE, bold=False)
    set_paragraph_borders(p, bottom=line_spec(COLOR_LINE, sz=6, space=2))
    return p


def add_summary_block(doc, abstract_text, meta_rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_fixed_layout(table)
    set_table_borders(
        table,
        top=line_spec(COLOR_LINE, sz=6, space=0),
        bottom=line_spec(COLOR_LINE, sz=6, space=0),
        left={"val": "nil"},
        right={"val": "nil"},
        insideH={"val": "nil"},
        insideV={"val": "nil"},
    )

    row = table.rows[0]
    c_label, c_body = row.cells
    c_label.width = Cm(1.15)
    c_body.width = Cm(PAGE_TEXT_WIDTH_CM - 1.15)

    c_label.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    c_body.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(c_label, top=60, bottom=60, left=40, right=40)
    set_cell_margins(c_body, top=60, bottom=60, left=40, right=40)
    set_cell_borders(c_label, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    set_cell_borders(c_body, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})

    p_label = c_label.paragraphs[0]
    p_label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p_label, before=0, after=0, line=LINE_META)
    r_label = p_label.add_run("摘要")
    set_run_font(r_label, FONT_SONG, size=Pt(9))

    p_abs = c_body.paragraphs[0]
    p_abs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.first_line_indent = Pt(21)
    set_para_spacing(p_abs, before=0, after=3, line=LINE_BODY)
    abs_content = split_inline_math(abstract_text)
    abs_has_math = add_rich_content(p_abs, abs_content, cn_font=FONT_SONG, en_font=FONT_EN, size=SIZE_ABSTRACT)
    if abs_has_math:
        p_abs.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        p_abs.paragraph_format.line_spacing = Pt(LINE_BODY)
    set_paragraph_borders(p_abs, bottom=line_spec(COLOR_LINE, sz=4, space=2))

    label_width_cm = 2.6
    for label, value in meta_rows:
        p = c_body.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_para_spacing(p, before=0, after=0, line=LINE_META)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(label_width_cm), WD_TAB_ALIGNMENT.LEFT)
        r1 = p.add_run(label)
        set_run_font(r1, FONT_SONG, size=SIZE_META)
        r2 = p.add_run("\t")
        set_run_font(r2, FONT_SONG, size=SIZE_META)
        value_has_math = add_rich_content(p, split_inline_math(value), cn_font=FONT_SONG, en_font=FONT_EN, size=SIZE_META)
        if value_has_math:
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            p.paragraph_format.line_spacing = Pt(LINE_META)

    return table


def add_toc_block(doc, items):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_fixed_layout(table)
    set_table_borders(
        table,
        top=line_spec(COLOR_LINE, sz=6, space=0),
        bottom=line_spec(COLOR_LINE, sz=6, space=0),
        left={"val": "nil"},
        right={"val": "nil"},
        insideH={"val": "nil"},
        insideV={"val": "nil"},
    )

    c_label, c_body = table.rows[0].cells
    c_label.width = Cm(1.15)
    c_body.width = Cm(PAGE_TEXT_WIDTH_CM - 1.15)
    c_label.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    c_body.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(c_label, top=40, bottom=40, left=40, right=40)
    set_cell_margins(c_body, top=40, bottom=40, left=100, right=40)
    set_cell_borders(c_label, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    set_cell_borders(c_body, left=line_spec(COLOR_LINE, sz=6, space=1), top={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"})

    p_label = c_label.paragraphs[0]
    p_label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p_label, before=0, after=0, line=LINE_META)
    r = p_label.add_run("目录")
    set_run_font(r, FONT_SONG, size=Pt(9))

    first = True
    if not items:
        items = ["正文"]

    for item in items:
        p = c_body.paragraphs[0] if first else c_body.add_paragraph()
        first = False
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Pt(6)
        set_para_spacing(p, before=0, after=0, line=LINE_META)
        r = p.add_run(item)
        set_run_font(r, FONT_SONG, size=SIZE_TOC)

    return table


def add_section_heading(doc, text, size=SIZE_SECTION, before=10, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_spacing(p, before=before, after=after, line=LINE_SECTION)
    set_paragraph_borders(
        p,
        left={"val": "single", "sz": 16, "space": 6, "color": COLOR_BLUE},
        bottom=line_spec(COLOR_LINE, sz=4, space=2),
    )
    r = p.add_run(text)
    set_run_font(r, FONT_HEI, size=size, bold=False)
    return p


def add_body_para(doc, text, indent_first=True, before=0, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_para_spacing(p, before=before, after=after, line=LINE_BODY)
    if indent_first:
        p.paragraph_format.first_line_indent = Pt(21)
    has_math = add_rich_content(p, text, cn_font=FONT_SONG, en_font=FONT_EN, size=SIZE_BODY)
    if has_math:
        # 行内公式使用至少行距，避免分式/上下标下沿被固定行距裁切。
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        p.paragraph_format.line_spacing = Pt(LINE_BODY)
    return p


def add_formula(doc, formula_text, number):
    # 用无边框双列表格模拟“公式居中、编号右对齐”
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_fixed_layout(table)
    set_table_borders(
        table,
        top={"val": "nil"},
        bottom={"val": "nil"},
        left={"val": "nil"},
        right={"val": "nil"},
        insideH={"val": "nil"},
        insideV={"val": "nil"},
    )
    c_formula, c_num = table.rows[0].cells
    c_formula.width = Cm(15.3)
    c_num.width = Cm(PAGE_TEXT_WIDTH_CM - 15.3)
    c_formula.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    c_num.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    # 公式在 Word 中通常高于正文字高，增加上下边距避免下沿被裁切。
    set_cell_margins(c_formula, top=90, bottom=90, left=0, right=0)
    set_cell_margins(c_num, top=90, bottom=90, left=0, right=0)
    set_cell_borders(c_formula, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    set_cell_borders(c_num, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})

    p1 = c_formula.paragraphs[0]
    p1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p1, before=0, after=0, line=None)
    p1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    if isinstance(formula_text, tuple):
        fallback_text, latex_text = formula_text
    else:
        fallback_text, latex_text = formula_text, None
    if latex_text:
        latex_text = normalize_latex_for_word(latex_text)

    if math2docx is not None and latex_text:
        try:
            math2docx.add_math(p1, latex_text)
        except Exception:
            r1 = p1.add_run(fallback_text)
            set_run_font(r1, FONT_SONG, en_font=FONT_EN, size=SIZE_FORMULA, italic=True)
    else:
        r1 = p1.add_run(fallback_text)
        set_run_font(r1, FONT_SONG, en_font=FONT_EN, size=SIZE_FORMULA, italic=True)

    p2 = c_num.paragraphs[0]
    p2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_spacing(p2, before=0, after=0, line=None)
    p2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r2 = p2.add_run(f"（{number}）")
    set_run_font(r2, FONT_SONG, en_font=FONT_EN, size=SIZE_BODY)
    return table


def add_author_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_spacing(p, before=4, after=4, line=LINE_META)
    r = p.add_run(text)
    set_run_font(r, FONT_SONG, size=SIZE_AUTHOR)
    return p


def add_reference_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_spacing(p, before=0, after=0, line=LINE_REF)
    has_math = add_rich_content(p, text, cn_font=FONT_SONG, en_font=FONT_EN, size=SIZE_REF)
    if has_math:
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        p.paragraph_format.line_spacing = Pt(LINE_REF)
    return p


def strip_quotes(text: str) -> str:
    t = text.strip()
    if len(t) >= 2 and ((t[0] == '"' and t[-1] == '"') or (t[0] == "'" and t[-1] == "'")):
        return t[1:-1]
    return t


def parse_scalar(value: str) -> Any:
    v = value.strip()
    if not v:
        return ""
    if v.startswith("[") and v.endswith("]"):
        inner = v[1:-1].strip()
        if not inner:
            return []
        return [strip_quotes(item.strip()) for item in inner.split(",") if item.strip()]
    return strip_quotes(v)


def parse_front_matter(fm_text: str) -> dict[str, Any]:
    data: dict[str, Any] = {}
    current_list_key = None

    for raw_line in fm_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip() or line.lstrip().startswith("#"):
            continue

        m_item = FRONT_MATTER_LIST_RE.match(line)
        if m_item and current_list_key:
            data.setdefault(current_list_key, [])
            data[current_list_key].append(strip_quotes(m_item.group(1).strip()))
            continue

        m_key = FRONT_MATTER_ITEM_RE.match(line)
        if not m_key:
            continue

        key = m_key.group(1).strip()
        value = m_key.group(2)
        parsed = parse_scalar(value)
        if parsed == "":
            data[key] = []
            current_list_key = key
        else:
            data[key] = parsed
            current_list_key = key if isinstance(parsed, list) else None

    return data


def split_front_matter_and_body(text: str) -> tuple[dict[str, Any], str]:
    normalized = text.lstrip("\ufeff").replace("\r\n", "\n").replace("\r", "\n")
    if not normalized.startswith("---\n"):
        return {}, normalized

    end = normalized.find("\n---\n", 4)
    if end < 0:
        return {}, normalized

    fm_text = normalized[4:end]
    body = normalized[end + 5 :]
    return parse_front_matter(fm_text), body


def read_display_equation(lines: list[str], start: int) -> tuple[str, int]:
    line = lines[start].strip()

    # 单行 $$...$$
    if line.startswith("$$") and line.endswith("$$") and len(line) > 4:
        return line[2:-2].strip(), start + 1

    # 起始 $$ 在本行，结束 $$ 在后续行
    if line.startswith("$$"):
        prefix = line[2:]
    else:
        prefix = ""

    parts = []
    if prefix:
        parts.append(prefix)

    i = start + 1
    while i < len(lines):
        cur = lines[i]
        marker = cur.find("$$")
        if marker >= 0:
            parts.append(cur[:marker])
            return "\n".join(p.rstrip() for p in parts).strip(), i + 1
        parts.append(cur)
        i += 1

    return "\n".join(p.rstrip() for p in parts).strip(), i


def parse_section_blocks(lines: list[str]) -> list[Block]:
    blocks: list[Block] = []
    para_lines: list[str] = []
    i = 0

    def flush_paragraph():
        nonlocal para_lines
        if not para_lines:
            return
        text = "".join(line.strip() for line in para_lines).strip()
        if text:
            blocks.append(Block(kind="paragraph", content=text))
        para_lines = []

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if stripped.startswith("$$"):
            flush_paragraph()
            latex, next_i = read_display_equation(lines, i)
            if latex:
                blocks.append(Block(kind="equation", content=latex))
            i = next_i
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i < len(lines) and NUMBER_LINE_RE.match(lines[i].strip()):
                i += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)$", raw)
        if bullet:
            flush_paragraph()
            text = bullet.group(1).strip()
            if text:
                blocks.append(Block(kind="paragraph", content=text))
            i += 1
            continue

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        para_lines.append(raw)
        i += 1

    flush_paragraph()
    return blocks


def parse_markdown_sections(body: str) -> list[Section]:
    lines = body.split("\n")
    sections: list[Section] = []

    preamble: list[str] = []
    current_title: str | None = None
    current_lines: list[str] = []

    for line in lines:
        m_heading = HEADING_RE.match(line)
        if m_heading:
            if current_title is None:
                if preamble:
                    blocks = parse_section_blocks(preamble)
                    if blocks:
                        sections.append(Section(title="", blocks=blocks))
                    preamble = []
            else:
                sections.append(Section(title=current_title.strip(), blocks=parse_section_blocks(current_lines)))
            current_title = m_heading.group(1).strip()
            current_lines = []
            continue

        if current_title is None:
            preamble.append(line)
        else:
            current_lines.append(line)

    if current_title is None:
        blocks = parse_section_blocks(preamble)
        if blocks:
            sections.append(Section(title="", blocks=blocks))
    else:
        sections.append(Section(title=current_title.strip(), blocks=parse_section_blocks(current_lines)))

    return sections


def normalize_heading(title: str) -> str:
    t = title.strip()
    t = re.sub(r"^[0-9]+[\s\-._、．]*", "", t)
    t = t.replace(" ", "")
    return t


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "；".join(str(x).strip() for x in value if str(x).strip())
    return str(value).strip()


def as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(x).strip() for x in value if str(x).strip()]
    text = str(value).strip()
    return [text] if text else []


def get_meta(meta: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in meta:
            return meta[key]
    return None


def split_inline_math(text: str) -> list[Any]:
    """将段落拆成普通文本和 $...$ 行内公式片段。"""
    segments: list[Any] = []
    buf: list[str] = []
    i = 0
    n = len(text)

    while i < n:
        ch = text[i]

        if ch == "\\" and i + 1 < n and text[i + 1] == "$":
            buf.append("$")
            i += 2
            continue

        if ch != "$":
            buf.append(ch)
            i += 1
            continue

        j = i + 1
        latex_buf: list[str] = []
        found = False
        while j < n:
            cj = text[j]
            if cj == "\\" and j + 1 < n and text[j + 1] == "$":
                latex_buf.append("$")
                j += 2
                continue
            if cj == "$":
                found = True
                break
            latex_buf.append(cj)
            j += 1

        if not found:
            buf.append(ch)
            i += 1
            continue

        if buf:
            segments.append("".join(buf))
            buf = []

        latex = "".join(latex_buf).strip()
        if latex:
            segments.append(m(latex, latex))
        i = j + 1

    if buf:
        segments.append("".join(buf))

    if not segments:
        return [""]
    return segments


def parse_entry_markdown(md_path: Path) -> EntryData:
    raw = md_path.read_text(encoding="utf-8")
    meta, body = split_front_matter_and_body(raw)
    sections = parse_markdown_sections(body)

    title = as_text(get_meta(meta, "title", "标题"))
    summary = as_text(get_meta(meta, "summary", "摘要"))
    pinyin = as_text(get_meta(meta, "pinyin", "拼音"))
    english_name = as_text(get_meta(meta, "english_name", "英文名称"))
    alias = as_text(get_meta(meta, "alias", "又称"))
    discipline = as_text(get_meta(meta, "discipline", "所属学科"))
    author = as_text(get_meta(meta, "author", "作者"))

    body_sections: list[Section] = []
    references: list[str] = []

    for section in sections:
        normalized = normalize_heading(section.title)
        if normalized == "摘要":
            if not summary:
                summary = "".join(block.content for block in section.blocks if block.kind == "paragraph")
            continue

        if normalized in {"参考文献", "参考资料"}:
            for block in section.blocks:
                if block.content.strip():
                    references.append(block.content.strip())
            continue

        if section.title or section.blocks:
            body_sections.append(section)

    if not title:
        for sec in body_sections:
            if sec.title:
                title = sec.title
                break
        if not title:
            title = md_path.stem

    if body_sections:
        first = body_sections[0]
        if first.title and normalize_heading(first.title) == normalize_heading(title) and not first.blocks:
            body_sections = body_sections[1:]

    if not summary:
        for sec in body_sections:
            for block in sec.blocks:
                if block.kind == "paragraph" and block.content.strip():
                    summary = block.content.strip()
                    break
            if summary:
                break

    toc = as_list(get_meta(meta, "toc", "目录"))
    if not toc:
        toc = [sec.title for sec in body_sections if sec.title]

    return EntryData(
        title=title,
        summary=summary,
        pinyin=pinyin,
        english_name=english_name,
        alias=alias,
        discipline=discipline,
        author=author,
        toc=toc,
        sections=body_sections,
        references=references,
    )


def setup_page(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)
    sec.top_margin = Cm(2.8)
    sec.bottom_margin = Cm(2.5)


def setup_default_styles(doc: Document):
    # 让 Word 原生公式等内容跟随更大的默认字号，而不仅仅依赖逐 run 设置。
    style = doc.styles["Normal"]
    style.font.name = FONT_EN
    style.font.size = SIZE_BODY

    rpr = style._element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        style._element.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_SONG)
    rfonts.set(qn("w:ascii"), FONT_EN)
    rfonts.set(qn("w:hAnsi"), FONT_EN)


def format_author(author: str) -> str:
    a = author.strip()
    if not a:
        return ""
    if a.startswith("（") or a.startswith("("):
        return a
    if a.startswith("作者"):
        return f"（{a}）"
    return f"（作者：{a}）"


def render_entry(entry: EntryData, output_path: Path):
    doc = Document()
    setup_page(doc)
    setup_default_styles(doc)

    add_entry_title(doc, entry.title)

    meta_rows = []
    if entry.pinyin:
        meta_rows.append(("拼　音", entry.pinyin))
    if entry.english_name:
        meta_rows.append(("英文名称", entry.english_name))
    if entry.alias:
        meta_rows.append(("又　称", entry.alias))
    if entry.discipline:
        meta_rows.append(("所属学科", entry.discipline))

    add_summary_block(doc, entry.summary, meta_rows)
    add_toc_block(doc, entry.toc)

    gap = doc.add_paragraph()
    set_para_spacing(gap, before=0, after=0, line=12)

    equation_no = 1
    for section in entry.sections:
        if section.title:
            add_section_heading(doc, section.title)

        for block in section.blocks:
            if not block.content.strip():
                continue
            if block.kind == "equation":
                add_formula(doc, (block.content, block.content), equation_no)
                equation_no += 1
                continue

            rich = split_inline_math(block.content)
            if len(rich) == 1 and isinstance(rich[0], str):
                add_body_para(doc, rich[0])
            else:
                add_body_para(doc, rich)

    author_text = format_author(entry.author)
    if author_text:
        add_author_line(doc, author_text)

    if entry.references:
        add_section_heading(doc, "参考文献", size=SIZE_SECTION_REF, before=12, after=6)
        for ref in entry.references:
            rich = split_inline_math(ref)
            if len(rich) == 1 and isinstance(rich[0], str):
                add_reference_para(doc, rich[0])
            else:
                add_reference_para(doc, rich)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="将百科条目 Markdown 生成为 Word 文档")
    parser.add_argument("-f", "--file", required=True, help="输入 Markdown 文件路径")
    parser.add_argument("-o", "--output", help="输出 Word 文件路径（默认输出到当前文件夹）")
    return parser.parse_args()


def main():
    args = parse_args()
    md_path = Path(args.file).expanduser().resolve()
    if not md_path.exists():
        raise SystemExit(f"未找到 Markdown 文件：{md_path}")

    if math2docx is None:
        print("警告：未安装 math2docx，数学公式将回退为普通文本。", file=sys.stderr)

    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else (Path.cwd() / f"{md_path.stem}.docx").resolve()
    )

    entry = parse_entry_markdown(md_path)
    render_entry(entry, output_path)
    print(f"文档已生成：{output_path}")


if __name__ == "__main__":
    main()
